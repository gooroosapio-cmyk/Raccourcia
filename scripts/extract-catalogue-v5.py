#!/usr/bin/env python3
"""Extrait le catalogue V5 du classeur vers data/catalogue/v5/*.json.

Le classeur est la source editoriale ; le depot n'en garde que ce qui sert a
l'import, en JSON, pour que le generateur de lots n'ait jamais a ouvrir un
fichier binaire et pour qu'une revue de code voie les donnees.

    python3 scripts/extract-catalogue-v5.py <classeur.xlsx>
    npx prettier --write data/catalogue/v5

La seconde ligne n'est pas facultative : ces fichiers sont relus en revue de
code, et le depot verifie leur mise en forme comme celle du reste.

Ce que ce script ne fait pas : il ne decide de rien. Les regroupements, les
renommages et les alias sont lus tels quels dans l'onglet Migration, et les
incoherences sont signalees plutot que corrigees en silence.
"""

import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

RACINE = Path(__file__).resolve().parent.parent
SORTIE = RACINE / 'data' / 'catalogue' / 'v5'


def lire(classeur, onglet):
    feuille = classeur[onglet]
    lignes = feuille.iter_rows(values_only=True)
    entetes = next(lignes)
    if not entetes or all(c is None for c in entetes):
        entetes = next(lignes)
    return [
        dict(zip(entetes, ligne)) for ligne in lignes if any(c is not None for c in ligne)
    ]


def liste(valeur):
    """« a; b; c » -> ["a", "b", "c"]."""
    if not valeur:
        return []
    return [x.strip() for x in str(valeur).split(';') if x.strip()]


def niveau(valeur):
    """« B — Direct guide » -> « B ». L'enum en base ne connait que la lettre."""
    lettre = str(valeur or '').strip()[:1].upper()
    if lettre not in ('A', 'B', 'C', 'D', 'E'):
        raise SystemExit(f'Niveau d\'execution illisible : {valeur!r}')
    return lettre


# L'enum `risk_level` date de la premiere migration et ne connait que trois
# valeurs sans accent. Le classeur V5 melange francais, anglais et accents :
# on normalise ici, sinon deux orthographes d'un meme niveau finiraient en
# filtres qui ne trouvent rien.
RISQUES = {
    'faible': 'faible', 'low': 'faible',
    'moyen': 'moyen', 'medium': 'moyen', 'modéré': 'moyen', 'modere': 'moyen',
    'élevé': 'eleve', 'eleve': 'eleve', 'high': 'eleve',
}


def risque(valeur):
    cle = str(valeur or '').strip().lower()
    if cle not in RISQUES:
        raise SystemExit(f'Niveau de risque inconnu : {valeur!r}')
    return RISQUES[cle]


def main():
    if len(sys.argv) != 2:
        raise SystemExit('Usage : extract-catalogue-v5.py <classeur.xlsx>')

    classeur = load_workbook(sys.argv[1], read_only=True, data_only=True)

    taxonomie = lire(classeur, 'Taxonomie')
    commandes = lire(classeur, 'Import_Commandes')
    payloads = lire(classeur, 'Payloads')
    questions = lire(classeur, 'Questions')
    migration = lire(classeur, 'Migration')

    # --- Familles ------------------------------------------------------
    familles = [
        {
            'family_id': f['family_id'],
            'domain': f['domain'],
            'name': f['name'],
            'short_description': f['short_description'],
            # Assemble a partir des colonnes du classeur, sans rien inventer :
            # metier, exigence de qualite, public, entrees, sorties.
            'description_long': (
                f"{f['expertise']}. {f['quality_rules']} "
                f"Pour : {f['users']}. "
                f"Formats acceptés : {f['input_formats']}. "
                f"Résultats : {f['output_types']}."
            ),
            'sort_order': int(f['order']),
        }
        for f in taxonomie
    ]

    # --- Commandes canoniques ------------------------------------------
    # Les colonnes medias ne sont pas extraites du tout : le classeur ne doit
    # pas pouvoir toucher un visuel envoye depuis l'administration.
    canoniques = []
    for c in commandes:
        presets = json.loads(c['presets_json'])
        canoniques.append({
            'ref': c['id'],
            'command': c['command'],
            'domain': c['domain'],
            'family_id': c['family_id'],
            'title': c['title'],
            'short_description': c['short_description'],
            'main_use_case': c['main_use_case'],
            'level': niveau(c['level']),
            'preset_key': presets[0] if presets else None,
            'input_primary': c['input_primary'],
            'required_variables': liste(c['required_variables']),
            'optional_variables': liste(c['optional_variables']),
            'sufficient_context': c['sufficient_context'],
            'blocking_condition': c['blocking_condition'],
            'default_values': c['default_values'],
            'preserve': c['preserve'],
            'avoid': c['avoid'],
            'output_format': c['output_format'],
            'risk_level': risque(c['risk_level']),
            'quality_score': int(c['quality_score']),
            'content_version': c['content_version'],
        })

    # --- Payloads -------------------------------------------------------
    # L'empreinte du classeur est reverifiee ici : un payload altere entre
    # l'edition et l'extraction ne doit pas entrer dans le depot.
    moteurs = {'ChatGPT': 'chatgpt', 'Claude': 'claude', 'Gemini': 'gemini'}
    textes = []
    for p in payloads:
        empreinte = hashlib.sha256(p['payload'].encode('utf-8')).hexdigest()
        if empreinte != p['sha256']:
            raise SystemExit(f"Empreinte incorrecte pour {p['id']} / {p['engine']}.")
        textes.append({
            'ref': p['id'],
            'moteur': moteurs[p['engine']],
            'payload': p['payload'],
            'sha256': empreinte,
        })

    # --- Questions ------------------------------------------------------
    # Le modele V5 pose une question a la fois et reevalue apres la reponse :
    # il n'y a plus de variable a substituer, donc plus de choix fermes.
    # `condition` dit quand poser, `delegation` ce que fait « Choisis pour
    # moi » ; les deux servent au moteur, pas a un formulaire.
    interrogations = [
        {
            'ref': q['id'],
            'sort_order': int(q['order']),
            'question': q['question'],
            'condition': q['condition'],
            'delegation': q['delegation'],
        }
        for q in questions
    ]

    # --- Alias et renommages --------------------------------------------
    # Une ligne « fusionnee » dont l'identifiant est aussi celui de la
    # destination n'est pas un alias : c'est la commande elle-meme qui change
    # de nom et elargit sa mission. La distinguer ici evite de creer un alias
    # qui pointerait sur lui-meme.
    alias, renommages = [], []
    for m in migration:
        if not str(m['decision']).startswith('fusion'):
            continue
        entree = {
            'ref': m['historical_id'],
            'command': m['historical_command'],
            'canonical_ref': m['canonical_id'],
            'canonical_command': m['canonical_command'],
            'preset': json.loads(m['preset_json']),
        }
        (renommages if m['historical_id'] == m['canonical_id'] else alias).append(entree)

    # --- Noms auxquels chaque commande repond ---------------------------
    # Quelqu'un qui a garde « /adsocial » ou « /emailpro » en tete doit
    # retrouver la commande qui fait le travail. On rassemble donc, pour
    # chaque commande canonique : ses modes, les noms des raccourcis devenus
    # ses modes, leurs titres d'origine, et son propre nom d'avant quand elle
    # a ete renommee.
    noms = {}

    def ajouter(ref, *valeurs):
        cible = noms.setdefault(ref, [])
        for valeur in valeurs:
            valeur = str(valeur or '').strip().lstrip('/')
            if valeur and valeur not in cible:
                cible.append(valeur)

    for c in commandes:
        ajouter(c['id'], *json.loads(c['presets_json']))
    for entree in alias + renommages:
        ajouter(
            entree['canonical_ref'],
            entree['command'],
            entree['preset'].get('mode'),
            entree['preset'].get('historical_title'),
        )

    for c in canoniques:
        c['aliases'] = noms.get(c['ref'], [])

    # Cible de rangement, mise de cote : la bascule de navigation s'en sert
    # bien apres l'import, quand les familles V5 deviennent visibles.
    rangement = [{'ref': c['ref'], 'family_id': c['family_id']} for c in canoniques]

    SORTIE.mkdir(parents=True, exist_ok=True)
    for nom, donnees in (
        ('familles', familles),
        ('commandes', canoniques),
        ('payloads', textes),
        ('questions', interrogations),
        ('alias', alias),
        ('renommages', renommages),
        ('rangement', rangement),
    ):
        chemin = SORTIE / f'{nom}.json'
        chemin.write_text(
            json.dumps(donnees, ensure_ascii=False, indent=2) + '\n', encoding='utf-8'
        )
        print(f'{nom:12s} {len(donnees):5d} lignes')


if __name__ == '__main__':
    main()
