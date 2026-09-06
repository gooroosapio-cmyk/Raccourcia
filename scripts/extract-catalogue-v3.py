#!/usr/bin/env python3
"""
Extrait le catalogue Application V2 du XLSX vers data/catalogue/v3/*.json.

Le classeur est la source de verite editoriale ; une fois importe, Supabase
redevient la seule source de verite runtime et le classeur un export d'archive.

L'extraction s'arrete si un controle bloquant de l'onglet Audit_Controles
echoue, ou si un invariant recalcule ici ne tient pas : le classeur peut se
declarer conforme et ne pas l'etre, on ne lui fait pas confiance sur parole.

    pip install openpyxl
    python3 scripts/extract-catalogue-v3.py
"""

import json
import pathlib
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl est requis : pip install openpyxl")

ROOT = pathlib.Path(__file__).resolve().parent.parent
SOURCE = ROOT / "data" / "source" / "RaccourcIA_Catalogue_Application_V2.xlsx"
OUT = ROOT / "data" / "catalogue" / "v3"

# Feuille source -> fichier JSON produit.
EXPORTS = {
    "familles": "Familles",
    "prompts": "Catalogue_App",
    "qcm": "Questions_QCM",
    "medias": "Medias_Defaut",
    "compatibilite": "Compatibilite_IA",
    "tests": "Tests_QA",
    "audit": "Audit_Controles",
    "regles": "Regles_Moteur",
}

# Regle R02 : une commande est globalement unique et minuscule.
MOTIF_COMMANDE = re.compile(r"^/[a-z][a-z0-9_]{2,31}$")


def sheet_to_records(worksheet):
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []
    entetes = [str(cell).strip() if cell is not None else "" for cell in rows[0]]

    records = []
    for row in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        record = {}
        for entete, cell in zip(entetes, row):
            if not entete:
                continue
            if cell is None:
                record[entete] = None
            elif isinstance(cell, bool):
                record[entete] = cell
            elif isinstance(cell, (int, float)):
                record[entete] = cell
            else:
                valeur = str(cell).strip()
                record[entete] = valeur or None
        records.append(record)
    return records


def echouer(message):
    sys.exit(f"Extraction interrompue : {message}")


def verifier(donnees):
    """
    Invariants recalcules sur les donnees, pas lus dans l'onglet d'audit.

    Un classeur qui se declare conforme peut ne pas l'etre ; ces controles
    sont ceux dont depend l'integrite de la base, et ils bloquent.
    """
    prompts = donnees["prompts"]
    familles = donnees["familles"]
    qcm = donnees["qcm"]

    if len(prompts) != 320:
        echouer(f"{len(prompts)} raccourcis au lieu de 320.")

    commandes = [p["command"] for p in prompts]
    if len(set(commandes)) != len(commandes):
        echouer("des commandes sont dupliquees.")

    identifiants = [p["id"] for p in prompts]
    if len(set(identifiants)) != len(identifiants):
        echouer("des identifiants sont dupliques.")

    hors_motif = [c for c in commandes if not MOTIF_COMMANDE.match(c or "")]
    if hors_motif:
        echouer(f"commandes hors motif R02 : {hors_motif[:5]}")

    ids_familles = {f["family_id"] for f in familles}
    orphelins = [p["command"] for p in prompts if p["family_id"] not in ids_familles]
    if orphelins:
        echouer(f"raccourcis sans famille : {orphelins[:5]}")

    image = [f for f in familles if f["domain"] == "IMAGE"]
    texte = [f for f in familles if f["domain"] == "TEXTE"]
    if len(image) != 6 or len(texte) != 7:
        echouer(f"{len(image)} categories IMAGE et {len(texte)} TEXTE au lieu de 6 et 7.")

    # R06 : un raccourci IMAGE doit produire une image, jamais une explication.
    manquants = [p["command"] for p in prompts if p["domain"] == "IMAGE" and p["output_type"] != "image"]
    if manquants:
        echouer(f"raccourcis IMAGE sans sortie image : {manquants[:5]}")

    # R05 : trois questions au maximum, sinon le QCM devient un formulaire.
    par_prompt = {}
    for q in qcm:
        par_prompt.setdefault(q["prompt_id"], []).append(q)
    trop = [pid for pid, items in par_prompt.items() if len(items) > 3]
    if trop:
        echouer(f"plus de 3 questions pour : {trop[:5]}")

    sans_payload = [p["command"] for p in prompts if not (p.get("payload_copy") or "").strip()]
    if sans_payload:
        echouer(f"payload vide pour : {sans_payload[:5]}")

    bloquants = [
        a for a in donnees["audit"]
        if str(a.get("severity", "")).lower() == "bloquant"
        and str(a.get("status", "")).upper() not in ("OK", "PASS")
    ]
    if bloquants:
        echouer(f"controles bloquants en echec : {[a['audit_id'] for a in bloquants]}")


def main():
    if not SOURCE.exists():
        echouer(f"classeur introuvable : {SOURCE}")

    classeur = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    donnees = {}

    for nom_fichier, nom_feuille in EXPORTS.items():
        if nom_feuille not in classeur.sheetnames:
            echouer(f"feuille absente : {nom_feuille}")
        donnees[nom_fichier] = sheet_to_records(classeur[nom_feuille])

    verifier(donnees)

    OUT.mkdir(parents=True, exist_ok=True)
    for nom_fichier, enregistrements in donnees.items():
        chemin = OUT / f"{nom_fichier}.json"
        chemin.write_text(
            json.dumps(enregistrements, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(f"  {chemin.relative_to(ROOT)} : {len(enregistrements)} lignes")

    print("\nInvariants verifies : 320 raccourcis, commandes uniques et conformes,")
    print("6 categories IMAGE + 7 TEXTE, sorties image completes, QCM <= 3.")


if __name__ == "__main__":
    main()
